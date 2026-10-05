#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Tests for xlsx_add_column.py declared dimension bounds.

Regression tests for issue #5557: added formula or total cells could lie
outside the declared worksheet extent because the tool only widened the
right column of a colon-separated dimension — never the row bounds or
the top-left corner — and ignored single-cell or absent dimensions.
Read-only consumers (openpyxl read_only=True, Excel) trust the declared
extent and returned none of the added cells.

Fixtures materialize the repository workbook template (comments
stripped) with a package-absolute worksheet relationship, run the real
CLI, and repack with xlsx_pack before asserting on the result.
"""

import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
ADD_SCRIPT = os.path.join(SCRIPTS_DIR, "xlsx_add_column.py")
PACK_SCRIPT = os.path.join(SCRIPTS_DIR, "xlsx_pack.py")
TEMPLATE_DIR = os.path.join(os.path.dirname(SCRIPTS_DIR), "templates",
                            "minimal_xlsx")

try:
    from openpyxl import load_workbook
except ImportError:  # openpyxl is optional; only the round trip needs it
    load_workbook = None


def sheet_xml(dimension: str | None, cells: list[tuple[str, str]]) -> str:
    """Build sheetData rows from (ref, value) numeric cells."""
    rows = {}
    for ref, value in cells:
        m = re.fullmatch(r"([A-Z]+)([0-9]+)", ref)
        rows.setdefault(int(m.group(2)), []).append(
            f'<c r="{ref}"><v>{value}</v></c>')
    body = "".join(
        f'<row r="{r}">' + "".join(cells_xml) + "</row>"
        for r, cells_xml in sorted(rows.items())
    )
    dim = f'<dimension ref="{dimension}"/>' if dimension else ""
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        f"{dim}<sheetData>{body}</sheetData></worksheet>"
    )


def materialize(root: str, dimension: str | None,
                cells: list[tuple[str, str]]) -> str:
    """Copy the repository template, strip comments, set a controlled sheet."""
    work = os.path.join(root, "work")
    shutil.copytree(TEMPLATE_DIR, work)
    for dirpath, _, files in os.walk(work):
        for fn in files:
            path = os.path.join(dirpath, fn)
            with open(path, encoding="utf-8") as f:
                text = f.read()
            with open(path, "w", encoding="utf-8") as f:
                f.write(re.sub(r"<!--.*?-->", "", text, flags=re.S))
    rels = os.path.join(work, "xl", "_rels", "workbook.xml.rels")
    with open(rels, encoding="utf-8") as f:
        text = f.read()
    with open(rels, "w", encoding="utf-8") as f:
        f.write(text.replace('Target="worksheets/sheet1.xml"',
                             'Target="/xl/worksheets/sheet1.xml"'))
    with open(os.path.join(work, "xl", "worksheets", "sheet1.xml"), "w",
              encoding="utf-8") as f:
        f.write(sheet_xml(dimension, cells))
    return work


def run_add(work: str, *extra: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, ADD_SCRIPT, work, "--col", "G", "--sheet", "Sheet1",
         *extra],
        capture_output=True, text=True)


def dimension_of(work: str) -> str | None:
    with open(os.path.join(work, "xl", "worksheets", "sheet1.xml"),
              encoding="utf-8") as f:
        ws = f.read()
    m = re.search(r'<dimension ref="([^"]+)"', ws)
    return m.group(1) if m else None


def packed_formulas(work: str, root: str) -> list[tuple[str, str]]:
    out = os.path.join(root, "packed.xlsx")
    subprocess.run([sys.executable, PACK_SCRIPT, work, out],
                   capture_output=True, text=True, check=True)
    wb = load_workbook(out, read_only=True)
    try:
        ws = wb["Sheet1"]
        return [(c.coordinate, c.value) for row in ws.iter_rows()
                for c in row if c.data_type == "f"]
    finally:
        wb.close()


BASE_CELLS = [("A1", "10"), ("F1", "60"), ("A2", "20"), ("F2", "1000")]


class TestAddColumnDimensionBounds(unittest.TestCase):
    def test_formula_rows_extend_declared_rows(self):
        """G3:G5 added to A1:F2 must declare A1:G5 (issue #5557)."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:F2", BASE_CELLS)
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G5")

    def test_total_row_extends_declared_rows(self):
        """A total cell below the declared extent must widen the bounds."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:F2", BASE_CELLS)
            result = run_add(work, "--total-row", "7",
                             "--total-formula", "=SUM(G1:G6)")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G7")

    def test_single_cell_dimension_is_extended(self):
        """A single-cell dimension ref must be expanded too."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1", BASE_CELLS)
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G5")

    def test_absent_dimension_is_inserted_in_schema_order(self):
        """A sheet without <dimension> gets one, before sheetData."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, None, BASE_CELLS)
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G5")
            with open(os.path.join(work, "xl", "worksheets", "sheet1.xml"),
                      encoding="utf-8") as f:
                ws = f.read()
            self.assertLess(ws.index("<dimension"), ws.index("<sheetData"))

    def test_sparse_existing_cells_set_the_used_rectangle(self):
        """Existing cells outside the declared extent are included."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:F2",
                               BASE_CELLS + [("C4", "7"), ("F5", "9")])
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G5")

    def test_top_left_corner_expands_to_actual_cells(self):
        """Cells above/left of the declared corner pull it back out."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "B2:F2", BASE_CELLS)
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:G5")

    def test_no_cell_addition_leaves_dimension_untouched(self):
        """Nothing added means the declared extent must not change."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:F2", BASE_CELLS)
            result = run_add(work)  # no header/formula/total
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:F2")

    def test_larger_existing_extent_is_preserved(self):
        """Control: a declared extent wider than the cells is kept."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:Z100", BASE_CELLS)
            result = run_add(work, "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(dimension_of(work), "A1:Z100")

    @unittest.skipIf(load_workbook is None, "openpyxl not installed")
    def test_read_only_round_trip_returns_added_formulas(self):
        """openpyxl read_only must see every added formula cell."""
        with tempfile.TemporaryDirectory() as root:
            work = materialize(root, "A1:F2", BASE_CELLS)
            result = run_add(work, "--header", "Pct",
                             "--formula", "=F{row}/$F$2",
                             "--formula-rows", "3:5")
            self.assertEqual(result.returncode, 0, result.stderr)
            formulas = packed_formulas(work, root)
            self.assertEqual(
                [ref for ref, _ in formulas],
                ["G3", "G4", "G5"],
            )
            self.assertEqual(formulas[0][1], "=F3/$F$2")


if __name__ == "__main__":
    unittest.main()
