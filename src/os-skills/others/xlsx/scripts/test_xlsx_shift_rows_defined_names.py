#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
"""Regression tests for definedNames row shifting in xlsx_shift_rows.py.

Covers the named-range limitation: direct A1 cells/ranges, whole rows/columns
and comma unions in xl/workbook.xml <definedNames> must follow the global
row-shift policy, while quoted sheet qualifiers (including embedded "!",
apostrophes and commas), localSheetId/hidden attributes, complex expressions,
external/3D references and malformed or out-of-grid ranges are preserved
byte-for-byte. Also covers the non-mutating preview mode and an end-to-end
packed workbook round trip through xlsx_unpack.py / xlsx_pack.py / openpyxl.
"""

import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile
import unittest
import xml.etree.ElementTree as ET

NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"

SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
SCRIPT = os.path.join(SCRIPTS_DIR, "xlsx_shift_rows.py")
UNPACK = os.path.join(SCRIPTS_DIR, "xlsx_unpack.py")
PACK = os.path.join(SCRIPTS_DIR, "xlsx_pack.py")

try:
    import openpyxl  # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


def _load_module():
    spec = importlib.util.spec_from_file_location("xlsx_shift_rows_mod", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


# Keys map to <definedName name="..."> for the shared fixture below. Values
# marked shift are the expected text after `insert 2 at row 5`; the rest must
# survive byte-for-byte.
FIXTURE_NAMES = {
    "PlainRange": ("Sheet1!$A$5:$D$20", "Sheet1!$A$7:$D$22"),
    "MixedCell": ("Sheet1!$B7", "Sheet1!$B9"),
    "LocalCell": ("B7", "B9"),
    "QuotedBang": ("'My!Sheet'!$B$7:$C$9", "'My!Sheet'!$B$9:$C$11"),
    "Apostrophes": ("'It''s Data'!A7", "'It''s Data'!A9"),
    "QuotedComma": ("'A, B'!$A$7", "'A, B'!$A$9"),
    "Union": (
        "Sheet1!$A$1:$B$4,Sheet1!$D$10:$E$12",
        "Sheet1!$A$1:$B$4,Sheet1!$D$12:$E$14",
    ),
    "PrintTitles": ("Sheet1!$A:$C,Sheet1!$4:$9", "Sheet1!$A:$C,Sheet1!$4:$11"),
    "PrintArea": ("Sheet1!$A$1:$D$12", "Sheet1!$A$1:$D$14"),
    "HiddenName": ("Sheet1!$C$5", "Sheet1!$C$7"),
    "ComplexExpr": ("SUM(Sheet1!A5:A9)*2", None),
    "ExternalRef": ("[1]Sheet1!$A$5", None),
    "ThreeD": ("Sheet1:Sheet3!$A$5", None),
    "Malformed": ("Sheet1!$B$5:$C", None),
    "OutOfGrid": ("Sheet1!A0", None),
    "TooFar": ("Sheet1!A1048577", None),
    "BadUnionMember": ("Sheet1!$A$1:$B$2,Summary!Total", None),
    "QuotedExternal": ("'[book.xlsx]Sheet1'!$A$5", None),
    "SpacedValue": ("Sheet1!$A$1 , Sheet1!$B$2", None),
}


def _workbook_xml(names_subset=None):
    entries = []
    for name, (old, _new) in FIXTURE_NAMES.items():
        if names_subset is not None and name not in names_subset:
            continue
        attrs = f'name="{name}"'
        if name in ("QuotedBang", "PrintTitles", "PrintArea"):
            attrs += ' localSheetId="0"'
        if name == "HiddenName":
            attrs += ' hidden="1"'
        entries.append(f"<definedName {attrs}>{old}</definedName>")
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<workbook xmlns="{NS_MAIN}">'
        "<sheets><sheet name=\"Sheet1\" sheetId=\"1\"/></sheets>"
        f"<definedNames>{''.join(entries)}</definedNames>"
        "</workbook>\n"
    )


SHEET1_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<worksheet xmlns="{NS_MAIN}"><sheetData>'
    '<row r="5"><c r="A5"><v>1</v></c></row>'
    "</sheetData></worksheet>\n"
)


def _make_work_dir(xml_text):
    work = tempfile.mkdtemp(prefix="shift_names_")
    os.makedirs(os.path.join(work, "xl", "worksheets"))
    with open(os.path.join(work, "xl", "workbook.xml"), "w", encoding="utf-8") as fh:
        fh.write(xml_text)
    with open(os.path.join(work, "xl", "worksheets", "sheet1.xml"), "w", encoding="utf-8") as fh:
        fh.write(SHEET1_XML)
    return work


def _run_script(work, op, at, count):
    return subprocess.run(
        [sys.executable, SCRIPT, work, op, str(at), str(count)],
        capture_output=True, text=True,
    )


def _defined_name_texts(work):
    tree = ET.parse(os.path.join(work, "xl", "workbook.xml"))
    texts = {}
    for dn in tree.getroot().iter(f"{{{NS_MAIN}}}definedName"):
        texts[dn.get("name")] = dn.text or ""
    return texts


class DefinedNamesShiftTests(unittest.TestCase):
    def setUp(self):
        self.work = _make_work_dir(_workbook_xml())
        self.addCleanup(shutil.rmtree, self.work, ignore_errors=True)

    def test_insert_shifts_cells_ranges_unions_and_print_areas(self):
        proc = _run_script(self.work, "insert", 5, 2)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        texts = _defined_name_texts(self.work)
        for name, (old, new) in FIXTURE_NAMES.items():
            if new is None:
                continue
            self.assertEqual(texts.get(name), new, f"{name}: {old!r} should become {new!r}")

        # Attributes must survive the rewrite untouched.
        tree = ET.parse(os.path.join(self.work, "xl", "workbook.xml"))
        by_name = {
            dn.get("name"): dn
            for dn in tree.getroot().iter(f"{{{NS_MAIN}}}definedName")
        }
        self.assertEqual(by_name["QuotedBang"].get("localSheetId"), "0")
        self.assertEqual(by_name["PrintArea"].get("localSheetId"), "0")
        self.assertEqual(by_name["HiddenName"].get("hidden"), "1")

    def test_delete_shifts_rows_up_with_floor_at_row_one(self):
        work = _make_work_dir(_workbook_xml({"PlainRange", "FloorRange"}).replace(
            "</definedNames>",
            '<definedName name="FloorRange">Sheet1!$B$2:$C$4</definedName></definedNames>',
        ))
        self.addCleanup(shutil.rmtree, work, ignore_errors=True)
        proc = _run_script(work, "delete", 1, 3)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        texts = _defined_name_texts(work)
        # at=1, delta=-3: rows 5,20 -> 2,17; rows 2,4 clamp at 1.
        self.assertEqual(texts["PlainRange"], "Sheet1!$A$2:$D$17")
        self.assertEqual(texts["FloorRange"], "Sheet1!$B$1:$C$1")

    def test_quoted_sheet_qualifiers_preserved_exactly(self):
        proc = _run_script(self.work, "insert", 5, 2)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        texts = _defined_name_texts(self.work)
        self.assertEqual(texts["QuotedBang"], "'My!Sheet'!$B$9:$C$11")
        self.assertEqual(texts["Apostrophes"], "'It''s Data'!A9")
        self.assertEqual(texts["QuotedComma"], "'A, B'!$A$9")

    def test_unsupported_values_preserved_byte_for_byte(self):
        proc = _run_script(self.work, "insert", 5, 2)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        texts = _defined_name_texts(self.work)
        for name, (old, new) in FIXTURE_NAMES.items():
            if new is not None:
                continue
            self.assertEqual(texts.get(name), old, f"{name} must be preserved exactly")

    def test_preview_mode_reports_plan_without_mutating_file(self):
        mod = _load_module()
        wb_path = os.path.join(self.work, "xl", "workbook.xml")
        with open(wb_path, "rb") as fh:
            before = fh.read()
        planned = mod.process_defined_names(wb_path, 5, 2, apply=False)
        with open(wb_path, "rb") as fh:
            after = fh.read()
        self.assertEqual(before, after, "preview must not rewrite the workbook part")
        got = {name: (old, new) for name, old, new in planned}
        self.assertEqual(got.get("PlainRange"), ("Sheet1!$A$5:$D$20", "Sheet1!$A$7:$D$22"))
        self.assertEqual(got.get("QuotedBang"), ("'My!Sheet'!$B$7:$C$9", "'My!Sheet'!$B$9:$C$11"))
        self.assertNotIn("ComplexExpr", got)
        self.assertNotIn("BadUnionMember", got)
        # The same preview plan applied for real must change the file.
        applied = mod.process_defined_names(wb_path, 5, 2, apply=True)
        self.assertEqual(applied, planned)
        with open(wb_path, "rb") as fh:
            self.assertNotEqual(before, fh.read())

    def test_cli_reports_defined_name_updates(self):
        proc = _run_script(self.work, "insert", 5, 2)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        self.assertIn("defined name", proc.stdout.lower())

    @unittest.skipUnless(HAVE_OPENPYXL, "openpyxl not installed")
    def test_packed_workbook_round_trip_aligns_cells_and_names(self):
        from openpyxl import Workbook, load_workbook
        from openpyxl.workbook.defined_name import DefinedName

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for row in range(1, 13):
            for col in range(1, 5):
                ws.cell(row=row, column=col, value=f"r{row}c{col}")
        ws.print_area = "A1:D12"
        ws.print_title_rows = "1:2"
        wb.defined_names["MyRange"] = DefinedName("MyRange", attr_text="Sheet1!$B$5:$C$9")

        src_xlsx = os.path.join(self.work, "src.xlsx")
        wb.save(src_xlsx)

        unpacked = os.path.join(self.work, "unpacked")
        subprocess.run([sys.executable, UNPACK, src_xlsx, unpacked], check=True,
                       capture_output=True, text=True)
        proc = _run_script(unpacked, "insert", 5, 2)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        out_xlsx = os.path.join(self.work, "out.xlsx")
        subprocess.run([sys.executable, PACK, unpacked, out_xlsx], check=True,
                       capture_output=True, text=True)

        rt = load_workbook(out_xlsx)
        ws2 = rt["Sheet1"]
        # Cells that were at rows >= 5 moved down by 2; rows above stayed.
        self.assertEqual(ws2["B4"].value, "r4c2")
        self.assertEqual(ws2["B5"].value, None)
        self.assertEqual(ws2["B7"].value, "r5c2")
        self.assertEqual(ws2["C11"].value, "r9c3")
        # Print area / titles / named range follow the same policy.
        self.assertEqual(ws2.print_area, "'Sheet1'!$A$1:$D$14")
        self.assertEqual(ws2.print_title_rows, "$1:$2")
        self.assertEqual(rt.defined_names["MyRange"].attr_text, "Sheet1!$B$7:$C$11")


if __name__ == "__main__":
    unittest.main()
